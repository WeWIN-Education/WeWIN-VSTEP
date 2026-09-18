from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "public" / "templates" / "WEWIN_VSTEP_Exam_Import_Template.docx"
BLUE = RGBColor(0x00, 0x4A, 0xAD)

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
styles["Normal"].font.size = Pt(10.5)
for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
    styles[name].font.name = "Arial"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles[name].font.color.rgb = RGBColor(0,0,0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Mẫu nhập đề VSTEP cho WEWIN")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Dùng trong Word hoặc Google Docs rồi xuất thành DOCX")
run.italic = True

doc.add_heading("Cách sử dụng", level=1)
for text in [
    "Tạo một bản sao của tài liệu này. Chỉ sửa phần nằm sau dấu hai chấm và nội dung trong khối TEXT hoặc INSTRUCTIONS.",
    "Giữ nguyên các thẻ trong ngoặc vuông, tên trường và mã id. Mỗi id phải duy nhất.",
    "Đặt tên file audio đúng như trường audio trong tài liệu. Khi nhập đề, chọn DOCX và tất cả file audio cùng lúc.",
    "Xem trước kết quả kiểm tra. Hệ thống chỉ cho xuất bản khi đủ nội dung, đáp án và cấu trúc VSTEP.",
]:
    doc.add_paragraph(text, style="List Number")

doc.add_heading("Quy tắc nội dung", level=1)
doc.add_paragraph("VSTEP đầy đủ gồm 35 câu Nghe, 40 câu Đọc, 2 bài Viết và 3 phần Nói. Đáp án nằm trong DOCX nhưng chỉ được lưu ở phía máy chủ. Audio câu hỏi Speaking là tùy chọn; nếu bỏ trống, trình duyệt sẽ đọc đề bằng giọng máy.")

def code_line(text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    r.font.size = Pt(9)
    r.bold = bold
    if text.startswith("[") and text.endswith("]"):
        r.font.color.rgb = BLUE
        r.bold = True
    return p

def question(qid, number):
    code_line("[QUESTION]", True)
    code_line(f"id: {qid}")
    code_line(f"number: {number}")
    code_line("prompt: [Điền câu hỏi]")
    for letter in "ABCD": code_line(f"{letter}: [Điền lựa chọn {letter}]")
    code_line("answer: [A/B/C/D]")
    code_line("explanation: [Điền giải thích ngắn cho đáp án đúng]")
    doc.add_paragraph()

doc.add_page_break()
doc.add_heading("Thông tin đề", level=1)
for line in ["[EXAM]","slug: vstep-test-XX","title: [Điền tên đề]","subtitle: Bài luyện VSTEP bốn kỹ năng","target: B1-C1","duration_minutes: 179","strict_vstep: true"]: code_line(line)

doc.add_page_break()
doc.add_heading("Listening", level=1)
code_line("[LISTENING]")
code_line("[INSTRUCTIONS]")
code_line("[Điền hướng dẫn chung cho phần Nghe]")
code_line("[/INSTRUCTIONS]")
number = 1
for part_index, count in enumerate([8,12,15], start=1):
    doc.add_heading(f"Listening Part {part_index}", level=2)
    for line in ["[LISTENING_PART]",f"id: listening-part-{part_index}",f"title: Part {part_index}",f"audio: listening-part-{part_index}.mp3","duration_seconds: 0","instructions: [Điền hướng dẫn cho phần này]"]: code_line(line)
    for _ in range(count):
        question(f"listening-{number}", number); number += 1
    if part_index < 3: doc.add_page_break()

doc.add_page_break()
doc.add_heading("Reading", level=1)
code_line("[READING]")
code_line("[INSTRUCTIONS]")
code_line("[Điền hướng dẫn chung cho phần Đọc]")
code_line("[/INSTRUCTIONS]")
number = 1
for passage_index in range(1,5):
    doc.add_heading(f"Reading Passage {passage_index}", level=2)
    for line in ["[READING_PASSAGE]",f"id: reading-passage-{passage_index}",f"title: Passage {passage_index}","[TEXT]","[Điền toàn bộ bài đọc. Có thể dùng nhiều đoạn văn.]","[/TEXT]"]: code_line(line)
    for _ in range(10):
        question(f"reading-{number}", number); number += 1
    if passage_index < 4: doc.add_page_break()

doc.add_page_break()
doc.add_heading("Writing", level=1)
code_line("[WRITING]")
for index, duration, words in [(1,20,120),(2,40,250)]:
    doc.add_heading(f"Writing Task {index}", level=2)
    for line in ["[WRITING_TASK]",f"id: writing-{index}",f"title: Task {index}",f"duration_minutes: {duration}",f"minimum_words: {words}","prompt: [Điền đề bài Writing]","bullet: [Điền yêu cầu; có thể lặp lại dòng bullet]"]: code_line(line)

doc.add_page_break()
doc.add_heading("Speaking", level=1)
code_line("[SPEAKING]")
for index, prep, speaking in [(1,15,180),(2,60,180),(3,60,240)]:
    doc.add_heading(f"Speaking Part {index}", level=2)
    for line in ["[SPEAKING_PART]",f"id: speaking-{index}",f"title: Part {index}",f"preparation_seconds: {prep}",f"speaking_seconds: {speaking}","audio:","prompt: [Điền đề bài Speaking]","question: [Điền câu hỏi phụ; có thể lặp lại dòng question]"]: code_line(line)

doc.core_properties.title = "Mẫu nhập đề VSTEP cho WEWIN"
doc.core_properties.subject = "Mẫu cấu trúc DOCX để nhập đề và audio tự động"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
